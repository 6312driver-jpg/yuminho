# -*- coding: utf-8 -*-
"""비트코인_멱법칙.xlsx 에 마지막 날짜 이후의 일봉(UTC, 마감된 날만)을 추가한다.

같은 폴더의 엑셀 파일을 갱신하며, 실행 기록은 update_btc.log 에 남는다.
데이터: Binance BTCUSDT 일봉 (실패 시 Coinbase BTC-USD 일봉).
"""
import json
import logging
import math
import sys
import urllib.request
from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Font

HERE = Path(__file__).resolve().parent
XLSX = HERE / '비트코인_멱법칙.xlsx'
LOG = HERE / 'update_btc.log'
FIT_ROWS = 8000  # '설정' 시트 회귀 범위와 동일
S = "'설정'!$B$"

logging.basicConfig(filename=LOG, level=logging.INFO, encoding='utf-8',
                    format='%(asctime)s %(levelname)s %(message)s')


def get_json(url):
    req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
    with urllib.request.urlopen(req, timeout=30) as r:
        return json.loads(r.read().decode())


def fetch_binance(start, end):
    """start~end(포함, UTC 날짜) 일봉 → {date: (open, close)}"""
    out = {}
    cur = start
    while cur <= end:
        ms = int(datetime(cur.year, cur.month, cur.day, tzinfo=timezone.utc).timestamp() * 1000)
        rows = get_json('https://api.binance.com/api/v3/klines?symbol=BTCUSDT'
                        f'&interval=1d&startTime={ms}&limit=1000')
        if not rows:
            break
        for k in rows:
            day = datetime.fromtimestamp(k[0] / 1000, timezone.utc).date()
            if start <= day <= end:
                out[day] = (float(k[1]), float(k[4]))
        cur = datetime.fromtimestamp(rows[-1][0] / 1000, timezone.utc).date() + timedelta(days=1)
    return out


def fetch_coinbase(start, end):
    out = {}
    cur = start
    while cur <= end:
        stop = min(cur + timedelta(days=299), end)
        rows = get_json('https://api.exchange.coinbase.com/products/BTC-USD/candles'
                        f'?granularity=86400&start={cur}T00:00:00Z&end={stop}T23:59:59Z')
        for t, _lo, _hi, o, c, _v in rows:
            day = datetime.fromtimestamp(t, timezone.utc).date()
            if start <= day <= end:
                out[day] = (float(o), float(c))
        cur = stop + timedelta(days=1)
    return out


def fetch(start, end):
    try:
        return fetch_binance(start, end)
    except Exception as e:
        logging.warning('Binance 실패 (%s) → Coinbase 사용', e)
        return fetch_coinbase(start, end)


def last_row(ws):
    r = ws.max_row
    while r > 1 and ws.cell(r, 1).value is None:
        r -= 1
    return r


def write_row(ws, i, day, o, c):
    ws.cell(i, 1, datetime(day.year, day.month, day.day)).number_format = 'yyyy-mm-dd'
    ws.cell(i, 2, round(o, 2))
    ws.cell(i, 3, round(c, 2))
    ws.cell(i, 4, f'=(L{i}-({S}5+{S}4*K{i}))/{S}6')
    ws.cell(i, 5, f'=H{i}*{S}8')
    ws.cell(i, 6, f'=H{i}*{S}9')
    ws.cell(i, 7, f'=A{i}-{S}3')
    ws.cell(i, 8, f'=10^({S}5+{S}4*K{i})')
    ws.cell(i, 9, f'=H{i}/{S}8')
    ws.cell(i, 10, f'=H{i}/{S}9')
    ws.cell(i, 11, f'=LOG10(G{i})')
    ws.cell(i, 12, f'=LOG10(C{i})')
    for j in range(1, 13):
        cell = ws.cell(i, j)
        cell.font = Font(name='Arial', color='0000FF' if j <= 3 else '000000')
        if j in (2, 3, 5, 6, 8, 9, 10):
            cell.number_format = '#,##0.00'
        elif j == 4:
            cell.number_format = '+0.00;-0.00;0.00'
        elif j == 7:
            cell.number_format = '#,##0'
        elif j >= 11:
            cell.number_format = '0.0000'


def rebuild_chart(wb, ws, last):
    # openpyxl 은 불러온 파일의 차트를 보존하지 못하므로 매번 다시 그린다
    if '차트' in wb.sheetnames:
        del wb['차트']
    cs = wb.create_sheet('차트')
    ch = LineChart()
    ch.title = '비트코인 멱법칙 밴드 (로그 스케일)'
    ch.height, ch.width = 14, 30
    ch.y_axis.scaling.logBase = 10
    ch.y_axis.title = 'USD'
    ch.y_axis.numFmt = '#,##0'
    ch.x_axis.number_format = 'yyyy'
    ch.x_axis.title = '날짜'
    for col in (3, 6, 5, 8, 9, 10):
        ch.add_data(Reference(ws, min_col=col, min_row=1, max_row=last), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=last))
    for s, colr in zip(ch.series, ['000000', 'C00000', 'F4B183', '70AD47', '9DC3E6', '2E75B6']):
        s.graphicalProperties.line.solidFill = colr
        s.graphicalProperties.line.width = 12000
        s.smooth = False
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    cs.add_chart(ch, 'A1')


def power_law(ws, genesis):
    """'설정' 시트와 같은 로그-로그 회귀 (SLOPE/INTERCEPT/STEYX) 를 파이썬으로 계산"""
    xs, ys = [], []
    for d, _o, c in ws.iter_rows(min_row=2, max_col=3, values_only=True):
        if d is not None and c:
            xs.append(math.log10((d - genesis).days))
            ys.append(math.log10(c))
    n = len(xs)
    mx, my = sum(xs) / n, sum(ys) / n
    b = sum((x - mx) * (y - my) for x, y in zip(xs, ys)) / sum((x - mx) ** 2 for x in xs)
    a = my - b * mx
    sigma = math.sqrt(sum((y - a - b * x) ** 2 for x, y in zip(xs, ys)) / (n - 2))
    return a, b, sigma


def summary(wb):
    ws = wb['멱법칙']
    genesis = wb['설정']['B3'].value
    a, b, sig = power_law(ws, genesis)
    r = last_row(ws)
    day, close = ws.cell(r, 1).value, ws.cell(r, 3).value
    log_mid = a + b * math.log10((day - genesis).days)
    z = (math.log10(close) - log_mid) / sig
    mid = 10 ** log_mid
    return (f'[BTC 멱법칙] {day:%Y-%m-%d}\n'
            f'종가 ${close:,.0f}\n'
            f'현재 시그마 {z:+.2f}\n'
            f'중심가 ${mid:,.0f}\n'
            f'+1σ ${mid * 10 ** sig:,.0f} / +2σ ${mid * 10 ** (2 * sig):,.0f}\n'
            f'-1σ ${mid / 10 ** sig:,.0f}')


def notify(wb):
    try:
        import kakao_notify
    except ImportError:
        return
    if not kakao_notify.configured():
        return
    try:
        msg = summary(wb)
        kakao_notify.send(msg)
        logging.info('카카오톡 전송: %s', msg.replace('\n', ' | '))
    except Exception as e:
        logging.error('카카오톡 전송 실패: %s', e)


def main(fetcher=fetch, today=None):
    wb = load_workbook(XLSX)
    ws = wb['멱법칙']
    r = last_row(ws)
    last_day = ws.cell(r, 1).value.date()
    # 오늘(UTC) 일봉은 아직 마감 전이므로 어제까지만 받는다 (한국시간 09:00 = UTC 00:00 마감)
    end = (today or datetime.now(timezone.utc).date()) - timedelta(days=1)
    start = last_day + timedelta(days=1)
    if start > end:
        logging.info('추가할 데이터 없음 (마지막 %s)', last_day)
        notify(wb)
        return 0
    data = fetcher(start, end)
    days = sorted(data)
    for k, day in enumerate(days, start=r + 1):
        write_row(ws, k, day, *data[day])
    new_last = r + len(days)
    if new_last + 1 > FIT_ROWS:
        logging.warning('데이터가 %d행에 가까워졌습니다. 설정 시트 회귀 범위를 늘리세요.', FIT_ROWS)
    ws.auto_filter.ref = f'A1:L{new_last}'
    ws.conditional_formatting = ConditionalFormattingList()
    ws.conditional_formatting.add(f'D2:D{new_last}', ColorScaleRule(
        start_type='num', start_value=-2, start_color='5B9BD5', mid_type='num', mid_value=0,
        mid_color='FFFFFF', end_type='num', end_value=2, end_color='F8696B'))
    rebuild_chart(wb, ws, new_last)
    wb.calculation.fullCalcOnLoad = True  # 엑셀에서 열 때 수식 재계산
    wb.save(XLSX)
    logging.info('%d일 추가 (%s ~ %s)', len(days), days[0] if days else '-', days[-1] if days else '-')
    notify(wb)
    return len(days)


if __name__ == '__main__':
    try:
        main()
    except PermissionError:
        logging.error('엑셀 파일이 열려 있어 저장하지 못했습니다. 파일을 닫고 다시 실행하세요.')
        sys.exit(1)
    except Exception:
        logging.exception('업데이트 실패')
        sys.exit(1)
